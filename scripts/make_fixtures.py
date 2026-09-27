# Genera los fixtures de prueba de las cargas a partir de los datos del demo (app.html).
#  - fixtures/contrato-grupo-nf-ingenio-providencia.docx / .pdf  (texto literal de contratoTexto)
#  - fixtures/manifiesto-01278713.pdf                            (datos del objeto manifiesto del demo; lo que el demo no trae queda vacío)
#  - fixtures/manifiesto-prueba-ficticio.pdf                     (datos FICTICIOS rotulados como tales, para probar las reglas de 8 h, 5 días hábiles, SICE-TAC y valor en letras)
# Uso: python scripts/make_fixtures.py
import json, pathlib, subprocess
import fitz, docx
from docx.shared import Pt

root = pathlib.Path(__file__).resolve().parent.parent
out = root / 'fixtures'
out.mkdir(exist_ok=True)
html = (root / 'app.html').read_text(encoding='utf-8')
js = html[html.index('/* ===== DATOS DEMO'):html.index('/* ===== Sidebar + router')]
res = subprocess.run(['node', '-e', js + '\nprocess.stdout.write(JSON.stringify({contratoTexto,manifiesto}))'],
                     capture_output=True, text=True, encoding='utf-8', check=True)
data = json.loads(res.stdout)
texto = data['contratoTexto']
m = data['manifiesto']

# --- Contrato DOCX
d = docx.Document()
st = d.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11.5)
for i, par in enumerate(texto.split('\n\n')):
    p = d.add_paragraph()
    r = p.add_run(par.strip())
    if i == 0:
        r.bold = True
        p.alignment = 1
d.add_paragraph('\n\n______________________________                ______________________________')
d.add_paragraph('GRUPO NF S.A.S.                                              INGENIO PROVIDENCIA S.A.')
d.save(out / 'contrato-grupo-nf-ingenio-providencia.docx')


# --- Contrato PDF (texto seleccionable)
def pdf_text(path, blocks):
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    y = 60
    for txt, size, bold in blocks:
        font = 'tibo' if bold else 'tiro'
        h = page.insert_textbox(fitz.Rect(60, y, 552, 760), txt, fontsize=size, fontname=font)
        if h < 0:
            page = doc.new_page(width=612, height=792)
            y = 60
            h = page.insert_textbox(fitz.Rect(60, y, 552, 760), txt, fontsize=size, fontname=font)
        y += (760 - y) - h + 10
    doc.save(path)


pdf_text(out / 'contrato-grupo-nf-ingenio-providencia.pdf',
         [(par.strip(), 13 if i == 0 else 11, i == 0) for i, par in enumerate(texto.split('\n\n'))])


# --- Manifiestos PDF (formato tipo RNDC)
def manifiesto_pdf(path, rows, titulo, nota=None):
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((60, 50), titulo, fontsize=13, fontname='hebo')
    page.insert_text((60, 66), 'Registro Nacional de Despachos de Carga - RNDC', fontsize=9, fontname='helv')
    y = 84
    if nota:
        page.draw_rect(fitz.Rect(56, y - 2, 556, y + 26), color=(0.8, 0.1, 0.1), width=1)
        page.insert_textbox(fitz.Rect(60, y, 552, y + 26), nota, fontsize=8.5, fontname='hebo', color=(0.7, 0.1, 0.1))
        y += 34
    for sec, items in rows:
        page.draw_rect(fitz.Rect(56, y, 556, y + 15), color=None, fill=(0.14, 0.29, 0.35))
        page.insert_text((60, y + 11), sec, fontsize=9, fontname='hebo', color=(1, 1, 1))
        y += 19
        for k, v in items:
            page.insert_text((62, y + 9), k + ':', fontsize=8.5, fontname='hebo')
            if v:
                page.insert_textbox(fitz.Rect(250, y, 552, y + 26), v, fontsize=8.5, fontname='helv')
            page.draw_line(fitz.Point(60, y + 13), fitz.Point(552, y + 13), color=(0.85, 0.85, 0.85), width=0.4)
            y += 15
        y += 5
    doc.save(path)


def cop(v):
    return '$ ' + '{:,.0f}'.format(v).replace(',', '.')


manifiesto_pdf(out / 'manifiesto-01278713.pdf', [
    ('MANIFIESTO', [('Manifiesto N.', m['numero']), ('Autorización', m['autorizacion']), ('Tipo de manifiesto', m['tipo']), ('Fecha de expedición', m['expedicion'])]),
    ('EMPRESA DE TRANSPORTE', [('Razón social', m['empresa']), ('NIT', m['empresaNit'])]),
    ('REMITENTE Y DESTINATARIO', [('Remitente', m['remitente']), ('Destinatario', m['destinatario']), ('Origen', m['origen']), ('Destino', m['destino'])]),
    ('VEHÍCULO Y CONDUCTOR', [('Vehículo', m['vehiculo']), ('Poseedor / tenedor', m['poseedor']), ('Conductor', m['conductor'])]),
    ('MERCANCÍA', [('Descripción', m['mercancia']), ('Peso', m['peso'])]),
    ('VALORES', [('Valor total del viaje', cop(m['valorTotal'])), ('Valor en letras', m['valorLetras']), ('Retención ICA', cop(m['ica'])),
                 ('Saldo por pagar', cop(m['saldo'])), ('Fecha de pago del saldo', m['fechaPago'])]),
    ('CARGUE Y DESCARGUE', [('Cargue pagado por', m['carguePor']), ('Descargue pagado por', ''), ('Llegada / salida cargue', ''), ('Llegada / salida descargue', '')]),
    ('SEGUROS', [('Póliza', m['seguro'])]),
], 'MANIFIESTO ELECTRÓNICO DE CARGA')

manifiesto_pdf(out / 'manifiesto-prueba-ficticio.pdf', [
    ('MANIFIESTO', [('Manifiesto N.', '99000123'), ('Autorización', '100000001'), ('Tipo de manifiesto', 'General'),
                    ('Fecha de expedición', '03/03/2026'), ('Fecha cumplido', '05/03/2026')]),
    ('EMPRESA DE TRANSPORTE', [('Razón social', 'Transportes de Prueba S.A.S.'), ('NIT', '900.000.000-0')]),
    ('REMITENTE Y DESTINATARIO', [('Remitente', 'Generador de Prueba S.A. · NIT 900.000.001-1 · Calle 1 # 2-3, Cali'),
                                  ('Destinatario', 'Destinatario de Prueba S.A. · NIT 900.000.002-2 · Carrera 4 # 5-6, Bogotá'),
                                  ('Origen', 'Cali, Valle del Cauca · Calle 1 # 2-3'), ('Destino', 'Bogotá, D.C. · Carrera 4 # 5-6')]),
    ('VEHÍCULO Y CONDUCTOR', [('Vehículo', 'Kenworth · Placa XYZ123 · Config. 3S3'),
                              ('Poseedor / tenedor', 'Poseedor de Prueba · C.C. 1.000.000 · Calle 7 # 8-9, Cali'),
                              ('Conductor', 'Conductor de Prueba · C.C. 2.000.000')]),
    ('MERCANCÍA', [('Descripción', 'Carga general paletizada'), ('Peso', '30.000 kg')]),
    ('VALORES', [('Valor total del viaje', '$ 5.000.000'), ('Valor en letras', 'CINCO MILLONES QUINIENTOS MIL PESOS M/CTE'),
                 ('Retención ICA', '$ 20.000'), ('Saldo por pagar', '$ 5.000.000'),
                 ('Manifestación', 'La empresa de transporte manifiesta adeudar el saldo al titular del manifiesto.'),
                 ('Fecha de pago del saldo', '20/03/2026'), ('Valor SICE-TAC', '$ 5.400.000')]),
    ('CARGUE Y DESCARGUE', [('Cargue pagado por', 'Remitente'), ('Descargue pagado por', 'Empresa de transporte'),
                            ('Cita / llegada / salida cargue', '03/03/2026 06:00 / 03/03/2026 06:10 / 03/03/2026 17:30'),
                            ('Cita / llegada / salida descargue', '05/03/2026 08:00 / 05/03/2026 07:50 / 05/03/2026 13:00')]),
    ('SEGUROS', [('Póliza', 'Aseguradora de Prueba · Póliza P-0001')]),
], 'MANIFIESTO ELECTRÓNICO DE CARGA',
    'DOCUMENTO DE PRUEBA CON DATOS FICTICIOS. No corresponde a ningún manifiesto real; sirve para probar las reglas de verificación.')
print(sorted(p.name for p in out.iterdir()))
